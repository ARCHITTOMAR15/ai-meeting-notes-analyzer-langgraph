
from pathlib import Path
from docx import Document as DocxDocument
from llama_index.core.schema import Document


class TranscriptLoader:

    @staticmethod
    def load_document(file_path: str):
        path = Path(file_path)
        suffix = path.suffix.lower()

        if suffix == ".txt":
            text = path.read_text(encoding="utf-8")

        elif suffix == ".docx":
            doc = DocxDocument(file_path)
            text = "\n".join(p.text for p in doc.paragraphs)

        elif suffix == ".pdf":
            # Import only when a PDF is uploaded
            try:
                from PyPDF2 import PdfReader
            except ModuleNotFoundError:
                try:
                    from pypdf import PdfReader
                except ModuleNotFoundError:
                    raise ImportError(
                        "PDF support is unavailable because neither PyPDF2 nor pypdf is installed."
                    )

            reader = PdfReader(file_path)
            text = "\n".join(page.extract_text() or "" for page in reader.pages)

        else:
            raise ValueError(f"Unsupported file format: {suffix}")

        return Document(text=text)


